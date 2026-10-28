"""Build the SRS and System Design as Word documents (then PDF via Word).

Usage:  python tools/build_docs.py
Writes docs/*.docx; tools/export_pdf.ps1 turns them into PDFs with Microsoft Word.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

UNIVERSITY = "Syrian Private University"
FACULTY = "Faculty of Dentistry"
PROJECT = "SPU Dental Clinic System"
AUTHORS = ["Hamza Almawla", "Zeannab Alhommsi"]

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"

DOCS_TO_BUILD = [
    ("SRS.md", "SPU-Dental-SRS.docx", "Software Requirements Specification"),
    ("system-design.md", "SPU-Dental-System-Design.docx", "System Design Document"),
]


# ---------- low-level helpers ----------

def set_cell_shading(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=100, right=100):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def set_table_borders(table, color="A6A6A6", size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_field(paragraph, instr, placeholder=""):
    """Insert a Word field (PAGE, NUMPAGES, TOC ...)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run._r.append(it)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def bottom_border(paragraph, color="1F3A5F", size=8):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    bdr.append(b)
    p_pr.append(bdr)


def set_run_fonts(run, name, cs="Arial"):
    run.font.name = name
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(attr), name)
    fonts.set(qn("w:cs"), cs)


# ---------- inline markdown ----------

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


def add_inline(paragraph, text, size=None, bold=False, color=None):
    text = text.replace("\\_", "_").replace("\\*", "*")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_fonts(run, MONO_FONT)
            run.font.size = Pt((size or 11) - 1.5)
            run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
            if bold:
                run.bold = True
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if bold:
            run.bold = True
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


# ---------- document setup ----------

def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    for name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 13, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        st = styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = ACCENT
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), BODY_FONT if attr != "w:cs" else "Arial")
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(theme_attr)) is not None:
                del fonts.attrib[qn(theme_attr)]
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(3)

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.3)


def cover_page(doc, title, date_text):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(UNIVERSITY.upper())
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(FACULTY)
    r.font.size = Pt(13)
    r.font.color.rgb = GREY

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PROJECT)
    r.font.size = Pt(18)
    r.font.color.rgb = GREY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = ACCENT
    bottom_border(p, size=12)

    for _ in range(8):
        doc.add_paragraph()
    rows = [
        ("Prepared by", ", ".join(AUTHORS)),
        ("Date", date_text),
        ("Status", "Draft for review"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.width, c1.width = Cm(4), Cm(8)
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = c0.paragraphs[0].add_run(k)
        r.bold = True
        r.font.color.rgb = ACCENT
        c1.paragraphs[0].add_run(v)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc_page(doc):
    p = doc.add_paragraph()
    r = p.add_run("Table of Contents")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(12)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
              "Right-click and choose Update Field to build the table of contents.")


def header_footer(section, title):
    hp = section.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run(f"{PROJECT} — {title}")
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    bottom_border(hp, color="BFBFBF", size=4)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for piece in ("Page ", "PAGE", " of ", "NUMPAGES"):
        if piece in ("PAGE", "NUMPAGES"):
            add_field(fp, piece, "1")
        else:
            fp.add_run(piece)
    for r in fp.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = GREY


# ---------- block markdown ----------

def parse_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def add_table(doc, rows, usable_width_cm):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_cell_margins(table)

    # column widths: each column fits its longest single word, the rest of the
    # page width is shared in proportion to content length
    word_cm, lengths = [], []
    for c in range(ncols):
        cells = [r[c] if c < len(r) else "" for r in rows]
        longest_word = max((len(w) for v in cells for w in v.replace("*", "").split()), default=4)
        word_cm.append(longest_word * 0.2 + 0.45)
        lengths.append(min(max(max(len(v) for v in cells), 6), 60))
    spare = max(usable_width_cm - sum(word_cm), 0)
    total = sum(lengths)
    widths = [Cm(word_cm[c] + spare * lengths[c] / total) for c in range(ncols)]

    font_size = 9.5 if ncols <= 4 else 9
    for i, row_vals in enumerate(rows):
        row = table.rows[i]
        no_split(row)
        for c in range(ncols):
            cell = row.cells[c]
            cell.width = widths[c]
            val = row_vals[c] if c < len(row_vals) else ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            # never leave a header row alone at the foot of a page; keep small tables whole
            if i == 0 or (len(rows) <= 8 and i < len(rows) - 1):
                para.paragraph_format.keep_with_next = True
            if i == 0:
                set_cell_shading(cell, "1F3A5F")
                add_inline(para, val, size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if i % 2 == 0:
                    set_cell_shading(cell, "F2F5F9")
                add_inline(para, val, size=font_size)
    repeat_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_code(doc, lines, usable_width_cm):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D9D9D9")
    set_cell_margins(table, 100, 100, 160, 160)
    cell = table.rows[0].cells[0]
    cell.width = Cm(usable_width_cm)
    set_cell_shading(cell, "F5F5F5")
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(lines):
        run = para.add_run(line)
        set_run_fonts(run, MONO_FONT, cs=MONO_FONT)
        run.font.size = Pt(8.5)
        if i < len(lines) - 1:
            run.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc, image_path, caption, usable_width_cm):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(image_path), width=Cm(usable_width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = GREY


def convert(md_path, out_path, title):
    lines = md_path.read_text(encoding="utf-8").splitlines()

    # first H1 and the "date · @author" line go on the cover page
    date_text = ""
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    while lines and not lines[0].strip():
        lines = lines[1:]
    if lines and "·" in lines[0]:
        date_text = lines[0].split("·")[0].strip()
        lines = lines[1:]

    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    usable = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # EMU -> cm
    cover_page(doc, title, date_text)
    toc_page(doc)

    # cover and contents stay without header/footer; the body section gets them,
    # with page numbers restarting at 1
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    header_footer(body, title)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    body._sectPr.append(pg)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            add_code(doc, code, usable)
            i += 1
            continue

        m = re.match(r"^!\[(.*)\]\((.+)\)$", stripped)
        if m:
            add_figure(doc, md_path.parent / m.group(2), m.group(1), usable)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = parse_row(lines[i])
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            add_table(doc, rows, usable)
            continue

        m = re.match(r"^(#{2,4})\s+(.*)", stripped)
        if m:
            level = len(m.group(1)) - 1
            h = doc.add_heading(level=level)
            add_inline(h, m.group(2))
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if m:
            # numbered list: write the number ourselves so each list restarts at 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            p.add_run(f"{m.group(1)}.\t")
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            add_inline(p, m.group(2))
            i += 1
            continue

        # paragraph (join soft-wrapped lines)
        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#|\||```|[-*]\s|\d+\.\s)", lines[i].strip()
        ):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        text = " ".join(buf)
        # a paragraph that is only a bold label acts as a small sub-heading
        if re.fullmatch(r"\*\*[^*]+\*\*", text):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            add_inline(p, text, color=ACCENT)
        else:
            add_inline(p, text)

    core = doc.core_properties
    core.title = f"{PROJECT} — {title}"
    core.author = ", ".join(AUTHORS)
    core.subject = f"{UNIVERSITY}, {FACULTY}"
    doc.save(out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    for src, out, title in DOCS_TO_BUILD:
        convert(DOCS / src, DOCS / out, title)
